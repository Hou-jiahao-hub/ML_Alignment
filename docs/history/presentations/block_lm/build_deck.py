"""Package the reviewed illustration with editable algorithm annotations.

Run with a Python environment providing python-pptx. No ACTS dependency.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'data/presentations/block_lm'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
prs.core_properties.title = 'Sensor block-LM：固定径迹，更新传感器位姿'

def text(slide, x, y, w, h, content, size=21, color='24313A', bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Microsoft YaHei'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
    return box

s = prs.slides.add_slide(prs.slide_layouts[6])
s.shapes.add_picture(str(OUT / 'block_lm_diagram.png'), 0, 0, width=prs.slide_width, height=prs.slide_height)
# Native PPT overlays correct the raster's small coordinate drift and make
# the two local diagrams and component-indexed objective editable.
def cover(x,y,w,h):
    q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    q.fill.solid(); q.fill.fore_color.rgb=RGBColor(255,255,255)
    q.line.fill.background()

def segment(x1,y1,x2,y2,color='24313A',width=1.5):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    q.line.color.rgb=RGBColor.from_string(color); q.line.width=Pt(width)

cover(9.72,1.64,6.04,5.98)
text(s,9.88,1.80,3.05,.34,'当前位姿：局部坐标',18,bold=True)
text(s,9.88,4.87,3.05,.34,'候选位姿：局部坐标',18,bold=True)
text(s,13.18,2.34,2.47,.4,'局部测量残差',17,bold=True)
text(s,13.18,2.92,2.47,.7,'rᵢ = mᵢ − πₛ(pᵢ)\npᵢ = hᵢ ∩ Sₛ(δₛ)',15)
for top,step,dx,dy in [(2.24,'k',.70,.66),(5.31,'trial',1.0,.88)]:
    cover(9.86,top,3.03,2.30)
    ox,oy=10.14,top+1.95
    segment(ox,oy,ox+2.52,oy); segment(ox,oy,ox,oy-1.74)
    text(s,ox+.03,top+.02,2,.3,'v = LOC1',13)
    text(s,ox+1.65,oy+.02,1,.3,'u = LOC0',12)
    mx,my=ox+1.72,oy-1.24
    px,py=ox+dx,oy-dy
    segment(px,py,mx,my,'7135A2',2)
    # Explicit arrowhead, pointing from prediction toward measurement.
    import math
    ang=math.atan2(my-py,mx-px)
    for side in [-.5,.5]:
        segment(mx,my,mx-.13*math.cos(ang+side),my-.13*math.sin(ang+side),'7135A2',2)
    for kind,x,y,fill in [(MSO_SHAPE.DIAMOND,mx,my,'7135A2'),(MSO_SHAPE.OVAL,px,py,'FFFFFF')]:
        q=s.shapes.add_shape(kind,Inches(x-.065),Inches(y-.065),Inches(.13),Inches(.13))
        q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
        q.line.color.rgb=RGBColor.from_string('24313A')
    text(s,mx+.10,my-.28,1,.4,'mᵢ 固定',12,color='7135A2')
    text(s,px-.32,py+.08,1.8,.35,'πₛ(p_pred)',12)
    text(s,px+.02,py-.43,.8,.35,'rᵢ',16,color='7135A2')
cover(13.13,3.79,2.57,1.55)
text(s,13.18,3.82,2.47,.4,'传感器的 χ²',17,bold=True)
text(s,13.18,4.30,2.47,.7,'χ²ₛ = Σ (rᵢₐ / σᵢₐ)²\n(i, a) ∈ Fₛ',15)
text(s,13.18,5.01,2.47,.3,'a：有效 LOC 分量',12)
s.notes_slide.notes_text_frame.text = '''图为算法示意，不是实验结果。三个面表示同一个 sensor 的参考、当前和候选位姿，不是三个不同 sensor。右侧预测点应理解为交点转换到该面的局部坐标后的 π_s(p_pred)。候选使用 k+1 记号仅为图示，接受后才成为下一迭代。单条 hit 的残差缩短仅为示例；实际接受条件作用于同一有效分量集合的总 chi2，并不要求每条 hit 单独改善。当前和候选面的测量菱形表示同一个固定局部测量在各自几何下的全局嵌入。Nominal target 面与黑点不进入 target block 的残差目标；非 target nominal 几何用于建立固定径迹的参考。依据 scripts/ml_alignment_autograd.py::_sensor_residual_data、run_block_lm_alignment。'''

s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = RGBColor.from_string('F8F9F6')
text(s,.6,.4,14.8,.8,'读图与算法：一个 sensor，就是一个几何参数块',32,bold=True)
items = [
    ('01  冻结径迹', '非 target 测量 refit 后，固定 helix 参数。\n同一 sensor 上的多条径迹共同约束该块。', 'DCEFE5'),
    ('02  更新几何并求交', '仅调整该 sensor 的 active DOF。\n每个候选重新求交，并转到候选面的局部坐标。', 'FFF0C4'),
    ('03  评分与接受', '在本次线性化固定的有效分量集上比较 χ²。\n下降、不丢分量、失效率检查通过，才接受。', 'EEE3F6'),
]
for i,(title,body,fill) in enumerate(items):
    x=.6+i*5.0
    rect=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x), Inches(1.6), Inches(4.8), Inches(2.5))
    rect.fill.solid(); rect.fill.fore_color.rgb=RGBColor.from_string(fill)
    rect.line.fill.background()
    text(s,x+.2,1.8,4.4,.5,title,24,bold=True)
    text(s,x+.2,2.5,4.4,1.4,body,18)
text(s,.7,4.3,7.2,.5,'图中符号如何理解',26,bold=True)
text(s,.7,5.0,7.2,2.5,'绿色 helix：整个 block-LM 阶段固定\n灰框 / 黑点：nominal 对照 / 参考交点\n黄框 / 白点：当前或候选面 / 预测交点\n紫色菱形 / 箭头：局部测量 / 测量减预测',21)
text(s,8.4,4.3,6.8,.5,'讲解时需要强调',26,bold=True)
text(s,8.4,5.0,6.8,2.8,'测量的局部坐标固定；其全局位置随面移动。\n最小化加权局部残差，而非黑白两点距离。\n总 χ² 下降不要求每一条 hit 的残差都缩小。\n停止迭代不等于通过驻点与秩亏检查。',21)
text(s,.7,8.2,14.6,.4,'来源：ml_alignment_autograd.py · _sensor_residual_data / _block_lm_linearized_block / run_block_lm_alignment',12,color='62706E')
s.notes_slide.notes_text_frame.text = '''分块精确解耦源于径迹固定且每块残差仅依赖自身 alignment_row。标准化残差 r/σ 的 Jacobian 按 physical_scales 归一化，形成 H=JᵀJ 和 g=Jᵀr，解正则化 LM 步。实现中的正则项为 [max(0, eigen_floor-min_eigenvalue)+damping*max(max_eigenvalue,1)]I。步长截断与硬参数边界不同。候选拒绝在同一块复用 Jacobian 提高阻尼重试；接受后阻尼减半。固定集合仅限本次线性化及重试，下一次访问可能扩大，因此 history 不保证单调。终态区分 intersection_failed、rank_deficient、converged、budget_exhausted。'''
# Suppress theme-provided shadows, including on opaque correction covers.
from pptx.oxml.xmlchemy import OxmlElement
for slide in prs.slides:
    for shape in slide.shapes:
        props = shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if props is not None:
            for old in list(props.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')):
                props.remove(old)
            props.append(OxmlElement('a:effectLst'))
prs.save(OUT / 'block_lm_explained.pptx')
print(OUT / 'block_lm_explained.pptx')
